# -*- coding: utf-8 -*-
# Copyright 2026 Yves Vallée
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).
"""Fiches de frais : numérotation, contrôles, Excel et justificatifs."""
import base64
import io
from datetime import date

from odoo.exceptions import UserError
from odoo.tests import common, tagged


def png(color=(200, 30, 30)):
    from PIL import Image
    output = io.BytesIO()
    Image.new('RGB', (60, 90), color).save(output, format='PNG')
    return output.getvalue()


def two_page_pdf():
    from reportlab.pdfgen import canvas
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    for text in ("page 1", "page 2"):
        pdf.drawString(100, 700, text)
        pdf.showPage()
    pdf.save()
    return output.getvalue()


@tagged('post_install', '-at_install')
class TestExpenseSheet(common.TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Sheet = cls.env['expense.scan.sheet']
        cls.employee = cls.env['hr.employee'].create({'name': "Léon Fiche"})
        Product = cls.env['product.product']
        cls.meal = Product.create({'name': "Repas fiche", 'can_be_expensed': True})
        cls.flat = Product.create({'name': "Forfait fiche", 'can_be_expensed': True,
                                   'standard_price': 36.4})
        cls.project = cls.env['project.project'].create({'name': "Mission fiche"})

    def expense(self, name, day, product=None, total=10.0, **values):
        return self.env['hr.expense'].create(dict({
            'name': name, 'employee_id': self.employee.id,
            'product_id': (product or self.meal).id, 'date': date(2026, 8, day),
            'total_amount_currency': total,
        }, **values))

    def attach(self, expense, raw, name, mimetype):
        return self.env['ir.attachment'].create({
            'name': name, 'raw': raw, 'mimetype': mimetype,
            'res_model': 'hr.expense', 'res_id': expense.id,
        })

    def test_chronological_numbering_and_sub_numbers(self):
        late = self.expense("Tard", 20)
        early = self.expense("Tôt", 3)
        self.attach(early, png(), "a.png", 'image/png')
        self.attach(early, png(), "b.png", 'image/png')
        self.attach(late, png(), "c.png", 'image/png')
        flat = self.expense("Forfait", 10, product=self.flat, quantity=2)
        lines = self.Sheet._lines(late | early | flat)
        self.assertEqual([line['description'] for line in lines], ["Tôt", "Forfait", "Tard"])
        self.assertEqual(lines[0]['n'], "1.1 à 1.2")
        self.assertEqual([label for label, _a in lines[0]['receipts']], ["1.1", "1.2"])
        self.assertTrue(lines[1]['flat_rate'])
        self.assertFalse(lines[1]['receipts'])
        self.assertEqual(lines[1]['quantite'], 2)
        self.assertEqual(lines[2]['n'], "3")

    def test_vat_summary_groups_by_rate(self):
        Tax = self.env['account.tax']

        def purchase_tax(rate):
            return Tax.search([
                ('company_id', '=', self.env.company.id), ('type_tax_use', '=', 'purchase'),
                ('amount_type', '=', 'percent'), ('amount', '=', rate)], limit=1) \
                or Tax.create({'name': "Achats %s %% essai" % rate, 'amount': rate,
                               'amount_type': 'percent', 'type_tax_use': 'purchase'})

        tax_10, tax_20 = purchase_tax(10.0), purchase_tax(20.0)
        a = self.expense("Dix", 1, total=11.0, tax_ids=[(6, 0, tax_10.ids)])
        a.scan_tax_amount = 1.0
        b = self.expense("Dix bis", 2, total=22.0, tax_ids=[(6, 0, tax_10.ids)])
        b.scan_tax_amount = 2.0
        c = self.expense("Mêlés", 3, total=12.0, tax_ids=[(6, 0, tax_20.ids)])
        c.write({'scan_tax_amount': 1.4,
                 'scan_detected_tax': "plusieurs taux, jusqu'à 20.0 % — 1.40"})
        d = self.expense("Sans", 4, total=5.0, tax_ids=[(5, 0, 0)])
        rows = {row['rate']: row for row in self.Sheet._vat_summary(
            self.Sheet._lines(a | b | c | d))}
        self.assertEqual(rows[0.1]['total_tva'], 3.0)
        self.assertEqual(rows[0.1]['count'], 2)
        self.assertAlmostEqual(rows[0.2]['total_tva'], 1.4)
        self.assertIn("mêlés", rows[0.2]['label'])
        self.assertEqual(rows[0.0]['total_ttc'], 5.0)

    def test_reinvoice_only_template_refuses_other_expenses(self):
        template = self.env['expense.scan.export.template'].create({
            'name': "Refacturable", 'reinvoice_only': True})
        ok = self.expense("Oui", 1, reinvoice_mode='project', project_id=self.project.id)
        other = self.expense("Non", 2, reinvoice_mode='none')
        self.Sheet._check_template(template, ok)
        with self.assertRaises(UserError):
            self.Sheet._check_template(template, ok | other)

    def workbook(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl absent")
        book = openpyxl.Workbook()
        sheet = book.active
        sheet['A1'] = "SUIVI"
        sheet['C4'] = "ancienne prestation"
        for row in range(8, 11):
            sheet['A%s' % row] = row - 7
            sheet['E%s' % row] = None
            sheet['L%s' % row] = "=IF(OR($E%s=\"\",$E%s=0),0,$E%s*K%s)" % (row, row, row, row)
        sheet['L8'] = 99  # une ancienne donnée saisie à la main
        sheet.row_dimensions[10].hidden = True  # comme dans le modèle réel
        sheet.row_dimensions[11].height = 24.75
        sheet['A11'] = "TOTAL"
        sheet['L11'] = "=SUM(L8:L10)"
        sheet['K2'] = "=L11*2"  # un total repris ailleurs dans la feuille
        output = io.BytesIO()
        book.save(output)
        template = self.env['expense.scan.export.template'].create({
            'name': "Essai", 'file': base64.b64encode(output.getvalue()),
            'first_row': 8, 'last_row': 10,
            'column_ids': [
                (0, 0, {'cell': 'A', 'value': 'n'}),
                (0, 0, {'cell': 'B', 'value': 'date'}),
                (0, 0, {'cell': 'L', 'value': 'total_ttc'}),
                (0, 0, {'cell': 'C4', 'value': 'prestation'}),
            ],
        })
        return openpyxl, template

    def test_excel_fills_and_clears(self):
        openpyxl, template = self.workbook()
        expense = self.expense("Seul", 5, total=12.5, reinvoice_mode='project',
                               project_id=self.project.id)
        # La mission n'est citée que si l'on a filtré dessus.
        content = self.Sheet.with_context(
            expense_scan_sheet_project_ids=self.project.ids)._excel(template, expense)
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        self.assertEqual(sheet['A8'].value, "1")
        self.assertEqual(sheet['L8'].value, 12.5)
        self.assertEqual(sheet['B8'].value.date(), date(2026, 8, 5))
        self.assertEqual(sheet['C4'].value, "Mission fiche")
        unfiltered = openpyxl.load_workbook(io.BytesIO(
            self.Sheet._excel(template, expense))).active
        self.assertEqual(unfiltered['C4'].value, "Édition des frais sélectionnés")
        # Les lignes vides ont disparu : le total suit immédiatement.
        self.assertEqual(sheet['A9'].value, "TOTAL")
        self.assertEqual(sheet['L9'].value, "=SUM(L8:L8)")
        self.assertEqual(sheet['K2'].value, "=L9*2")
        self.assertIsNone(sheet['A11'].value)

    def test_basic_template_is_delivered_and_usable(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl absent")
        template = self.env.ref('expense_scan.export_template_basic')
        template.action_expense_scan_generate_file()
        self.assertTrue(template.filename.endswith(".xlsx"))
        blank = openpyxl.load_workbook(io.BytesIO(base64.b64decode(template.file))).active
        self.assertEqual(blank['K7'].value, "Sous-total TTC")
        self.assertEqual(blank['K10'].value, "=SUM(K8:K9)")
        self.assertEqual(blank['A3'].value, "Salarié")

        expenses = self.expense("Un", 1, total=10.0) | self.expense("Deux", 2, total=5.0) \
            | self.expense("Trois", 3, total=2.5)
        sheet = openpyxl.load_workbook(io.BytesIO(self.Sheet._excel(template, expenses))).active
        self.assertEqual(sheet['K10'].value, 2.5)
        self.assertEqual(sheet['K11'].value, "=SUM(K8:K10)")
        self.assertEqual(sheet['B3'].value, self.employee.name)

    def test_excel_grows_beyond_its_rows(self):
        openpyxl, template = self.workbook()
        expenses = self.env['hr.expense']
        for day in range(1, 6):
            expenses |= self.expense("J%s" % day, day, total=float(day))
        content = self.Sheet._excel(template, expenses)
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        self.assertEqual(sheet['L12'].value, 5.0)
        self.assertEqual(sheet['A13'].value, "TOTAL")
        self.assertEqual(sheet['L13'].value, "=SUM(L8:L12)")
        self.assertEqual(sheet['K2'].value, "=L13*2")
        # Lignes ajoutées visibles, hauteur du total conservée.
        self.assertFalse(any(sheet.row_dimensions[row].hidden for row in range(8, 13)))
        self.assertEqual(sheet.row_dimensions[13].height, 24.75)

    def test_receipts_pdf_has_one_page_per_receipt(self):
        from odoo.tools.pdf import PdfFileReader
        first = self.expense("Photo", 1)
        second = self.expense("Facture", 2)
        self.attach(first, png(), "photo.png", 'image/png')
        self.attach(second, two_page_pdf(), "facture.pdf", 'application/pdf')
        flat = self.expense("Forfait", 3, product=self.flat)
        self.attach(flat, png(), "inutile.png", 'image/png')
        content = self.Sheet._receipts_pdf(self.Sheet._lines(first | second | flat))
        self.assertEqual(PdfFileReader(io.BytesIO(content), strict=False).getNumPages(), 3)

    def test_summary_renders(self):
        expense = self.expense("Récap", 4, total=20.0)
        html, _kind = self.env['ir.actions.report']._render_qweb_html(
            'expense_scan.action_report_expense_sheet', expense.ids,
            data={'expense_ids': expense.ids})
        self.assertIn("Récap", html.decode())

    def test_wizard_filters_and_zips_per_employee(self):
        other = self.env['hr.employee'].create({'name': "Jules Fiche"})
        mine = self.expense("A moi", 1, reinvoice_mode="project",
                            project_id=self.project.id)
        theirs = self.expense("A lui", 2, employee_id=other.id,
                              reinvoice_mode="project", project_id=self.project.id)
        skipped = self.expense("Pas refacturable", 3, reinvoice_mode="none")
        for expense in mine | theirs | skipped:
            self.attach(expense, png(), "t.png", "image/png")
        wizard = self.env["expense.scan.sheet.wizard"].with_context(
            active_model="hr.expense", active_ids=(mine | theirs | skipped).ids).create({
                "scope": "reinvoice", "summary": False, "receipts": True})
        self.assertEqual(wizard.selected_count, 2)
        wizard.action_generate()
        self.assertTrue(wizard.result_name.endswith(".zip"))

    def test_wizard_counts_while_editing(self):
        """Choisir une mission dans la fenêtre compte ses dépenses tout de suite."""
        from odoo.tests import Form
        on_mission = self.expense("Mission", 1, reinvoice_mode="project",
                                  project_id=self.project.id)
        elsewhere = self.expense("Ailleurs", 2, reinvoice_mode="project",
                                 project_id=self.env["project.project"].create(
                                     {"name": "Autre mission fiche"}).id)
        wizard = self.env["expense.scan.sheet.wizard"].with_context(
            active_model="hr.expense", active_ids=(on_mission | elsewhere).ids)
        with Form(wizard) as form:
            self.assertTrue(form.selected_summary.startswith("2 "))
            form.project_ids.add(self.project)
            self.assertTrue(form.selected_summary.startswith("1 "))

    def test_not_reinvoiced_expense_keeps_its_mission(self):
        expense = self.expense("Suivi", 1, reinvoice_mode="none")
        values = expense._expense_scan_project_values(self.project, reinvoice=False)
        self.assertEqual(values["reinvoice_mode"], "none")
        self.assertEqual(values["project_id"], self.project.id)
        self.assertNotIn("sale_order_id", values)

    def test_mission_change_moves_the_automatic_analytic(self):
        """L'imputation d'office suit la mission ; une répartition manuelle reste."""
        first = self.env["project.project"].create({"name": "Mission A analytique"})
        second = self.env["project.project"].create({"name": "Mission B analytique"})
        expense = self.expense("Imputée", 3, reinvoice_mode="none", project_id=first.id)
        if "account_id" not in first._fields or not first.account_id:
            self.skipTest("pas de plan analytique pour les missions")
        self.assertEqual(expense.analytic_distribution, {str(first.account_id.id): 100.0})

        expense.project_id = second
        self.assertTrue(second.account_id)
        self.assertEqual(expense.analytic_distribution, {str(second.account_id.id): 100.0})

        expense.project_id = False
        self.assertFalse(expense.analytic_distribution)

        manual = {str(first.account_id.id): 50.0, str(second.account_id.id): 50.0}
        expense.write({"project_id": first.id, "analytic_distribution": manual})
        expense.project_id = second
        self.assertEqual(expense.analytic_distribution, manual)

    def test_task_proposed_in_user_order(self):
        Task = self.env["project.task"]
        later = Task.create({"name": "Seconde", "project_id": self.project.id, "sequence": 20})
        first = Task.create({"name": "Première", "project_id": self.project.id, "sequence": 5})
        expense = self.expense("Tâche", 4)
        values = expense._expense_scan_project_values(self.project)
        self.assertEqual(values["expense_scan_task_id"], first.id)
        first.state = "1_done"
        values = expense._expense_scan_project_values(self.project)
        self.assertEqual(values["expense_scan_task_id"], later.id)
