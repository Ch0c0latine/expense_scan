# -*- coding: utf-8 -*-
# Copyright 2026 Yves Vallée
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl-3.0).
"""Fiches de frais : récapitulatif PDF, export Excel sur modèle, justificatifs.

Les trois sorties partagent la même numérotation : chronologique, de 1 à n
pour chaque fiche, et n.1 à n.x quand une dépense porte plusieurs
justificatifs. Les forfaits — kilométrage, barèmes — figurent dans les
tableaux mais n'ont pas de justificatif.

Une fiche par salarié : plusieurs salariés sélectionnés donnent plusieurs
fichiers, remis ensemble dans une archive.
"""
import base64
import io
import logging
import re
import zipfile
from copy import copy
from datetime import date

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools import format_date

_logger = logging.getLogger(__name__)

#: Valeurs qu'une colonne de ligne peut recevoir.
LINE_VALUES = [
    ('n', "N° de justificatif"),
    ('date', "Date"),
    ('categorie', "Catégorie"),
    ('description', "Description"),
    ('enseigne', "Enseigne"),
    ('mission', "Mission"),
    ('quantite', "Quantité (nuitées, km…)"),
    ('prix_unitaire_ht', "Montant unitaire HT"),
    ('total_ht', "Sous-total HT"),
    ('taux_tva', "Taux de TVA"),
    ('tva_unitaire', "TVA unitaire"),
    ('total_tva', "Montant de TVA"),
    ('ttc_unitaire', "Montant unitaire TTC"),
    ('total_ttc', "Sous-total TTC"),
    ('mode_paiement', "Payé par"),
    ('salarie', "Salarié"),
]
#: Valeurs qu'une cellule d'en-tête peut recevoir.
HEADER_VALUES = [
    ('salarie', "Salarié"),
    ('prestation', "Prestation (missions)"),
    ('mois', "Mois (date du premier jour)"),
    ('periode', "Période (texte)"),
    ('date_debut', "Date de début"),
    ('date_fin', "Date de fin"),
    ('total_ht', "Total HT"),
    ('total_tva', "Total TVA"),
    ('total_ttc', "Total TTC"),
    ('societe', "Société"),
]
#: Toutes les valeurs, sans doublon, pour le champ unique des correspondances.
VALUES = LINE_VALUES + [item for item in HEADER_VALUES
                        if item[0] not in {key for key, _label in LINE_VALUES}]
CELL_RE = re.compile(r"^([A-Z]{1,3})([0-9]*)$")
#: Référence (« $L$70 ») ou plage (« L8:L68 ») dans une formule de la même
#: feuille ; ni un nom de fonction (« LOG10( »), ni une autre feuille (« F2!A1 »).
CELL_REF_RE = re.compile(
    r"(?<![A-Za-z0-9_!'.$])(\$?[A-Z]{1,3}\$?)(\d+)"
    r"(?::(\$?[A-Z]{1,3}\$?)(\d+))?(?![0-9A-Za-z_(])")
#: Plus grand côté d'une photo de justificatif dans les PDF, en pixels :
#: environ 250 dpi sur une page A4, assez pour relire un ticket photographié.
IMAGE_MAX_SIDE = 2400
IMAGE_QUALITY = 88


class ExpenseScanExportTemplate(models.Model):
    _name = 'expense.scan.export.template'
    _description = "Modèle d'export Excel des frais"
    _order = 'sequence, name'

    name = fields.Char(string="Nom", required=True)
    sequence = fields.Integer(default=10)
    active = fields.Boolean(default=True)
    company_id = fields.Many2one('res.company', string="Société")
    file = fields.Binary(
        string="Fichier Excel", attachment=True,
        help="Le classeur tel qu'il est rempli à la main, données comprises : "
             "les lignes de dépenses sont vidées avant d'être remplies.")
    filename = fields.Char(string="Nom du fichier")
    sheet_name = fields.Char(
        string="Feuille", help="Vide : la première feuille du classeur.")
    first_row = fields.Integer(string="Première ligne de dépense", default=8, required=True)
    last_row = fields.Integer(
        string="Dernière ligne de dépense", default=68, required=True,
        help="Le tableau exporté compte exactement une ligne par dépense : "
             "les lignes en trop sont supprimées, celles qui manquent "
             "ajoutées, et les totaux placés dessous suivent.")
    reinvoice_only = fields.Boolean(
        string="Frais refacturables uniquement",
        help="Refuse l'export si l'une des dépenses n'est pas marquée "
             "« À refacturer : Oui ».")
    single_project = fields.Boolean(
        string="Une seule mission",
        help="Refuse l'export si les dépenses portent sur plusieurs missions.")
    column_ids = fields.One2many(
        'expense.scan.export.column', 'template_id', string="Correspondances")
    note = fields.Text(string="Notes")

    @api.constrains('first_row', 'last_row')
    def _check_rows(self):
        for template in self:
            if template.first_row < 1 or template.last_row < template.first_row:
                raise UserError(_("Les lignes de dépense du modèle « %s » sont incohérentes.",
                                  template.name))

    def action_expense_scan_generate_file(self):
        """Classeur vierge qui suit les correspondances du modèle.

        Point de départ à télécharger, mettre aux couleurs de l'entreprise
        puis redéposer : titres de colonnes, libellés d'en-tête, lignes de
        dépenses encadrées et totaux.
        """
        for template in self:
            content = template._expense_scan_blank_workbook()
            template.write({
                'file': base64.b64encode(content),
                'filename': "%s.xlsx" % re.sub(r'[\\/:*?"<>|]+', '-', template.name),
            })
        return True

    @api.model
    def _expense_scan_fill_blank_files(self):
        """Classeur vierge des modèles livrés qui n'en ont pas encore."""
        try:
            import openpyxl  # noqa: F401, PLC0415
        except ImportError:
            _logger.warning("openpyxl absent : modèles d'export livrés sans classeur")
            return
        templates = self.env.ref('expense_scan.export_template_basic',
                                 raise_if_not_found=False)
        if templates and not templates.file:
            templates.action_expense_scan_generate_file()

    def _expense_scan_blank_workbook(self):
        self.ensure_one()
        try:
            import openpyxl  # noqa: PLC0415
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: PLC0415
            from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: PLC0415
        except ImportError as error:
            raise UserError(_("La bibliothèque openpyxl manque sur le serveur.")) from error

        labels = dict(VALUES)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = (self.sheet_name or _("Frais"))[:31]
        thin = Side(style='thin', color='808080')
        box = Border(left=thin, right=thin, top=thin, bottom=thin)
        title_fill = PatternFill('solid', start_color='DDE7F0')
        bold = Font(bold=True)

        sheet['A1'] = _("Fiche de frais")
        sheet['A1'].font = Font(bold=True, size=14)

        # En-tête : chaque valeur, et son libellé dans la cellule de gauche.
        for column in self.column_ids.filtered(lambda c: c.kind == 'header' and c.value):
            cell = sheet[column.cell.strip().upper()]
            if cell.column > 1:
                label = sheet.cell(row=cell.row, column=cell.column - 1)
                label.value = dict(HEADER_VALUES).get(column.value, labels[column.value])
                label.font = bold
            cell.border = box
            if column.value in ('mois', 'date_debut', 'date_fin'):
                cell.number_format = 'DD/MM/YYYY'
            elif column.value in ('total_ht', 'total_tva', 'total_ttc'):
                cell.number_format = '#,##0.00'

        first, last = self.first_row, self.last_row
        lines = {column_index_from_string(c.cell.strip().upper()): c.value
                 for c in self.column_ids if c.kind == 'line' and c.value}
        formats = {'date': 'DD/MM/YYYY', 'taux_tva': '0%', 'quantite': '0.##'}
        amounts = {'prix_unitaire_ht', 'total_ht', 'tva_unitaire', 'total_tva',
                   'ttc_unitaire', 'total_ttc'}
        for index, value in lines.items():
            letter = get_column_letter(index)
            if first > 1:
                head = sheet.cell(row=first - 1, column=index, value=labels[value])
                head.font = bold
                head.fill = title_fill
                head.border = box
                head.alignment = Alignment(wrap_text=True, vertical='center')
            for row in range(first, last + 1):
                cell = sheet.cell(row=row, column=index)
                cell.border = box
                cell.number_format = '#,##0.00' if value in amounts \
                    else formats.get(value, 'General')
            total = sheet.cell(row=last + 1, column=index)
            total.font = bold
            total.border = box
            if value in amounts:
                total.value = "=SUM(%s%s:%s%s)" % (letter, first, letter, last)
                total.number_format = '#,##0.00'
            sheet.column_dimensions[letter].width = \
                30 if value in ('description', 'categorie', 'mission') else 14
        if lines:
            sheet.cell(row=last + 1, column=min(lines), value=_("TOTAL"))
        sheet.freeze_panes = sheet.cell(row=first, column=1)

        output = io.BytesIO()
        book.save(output)
        return output.getvalue()


class ExpenseScanExportColumn(models.Model):
    _name = 'expense.scan.export.column'
    _description = "Correspondance d'un modèle d'export"
    _order = 'kind desc, id'

    template_id = fields.Many2one(
        'expense.scan.export.template', required=True, ondelete='cascade')
    cell = fields.Char(
        string="Colonne ou cellule", required=True,
        help="Une lettre de colonne (« B ») : la valeur de chaque dépense, "
             "ligne par ligne. Une cellule (« K4 ») : une valeur d'en-tête, "
             "écrite une fois. Pour une cellule fusionnée, sa première "
             "cellule en haut à gauche.")
    # Le type se déduit de la cellule : une seule chose à saisir, et plus de
    # colonne de valeur qui reste fermée selon le type choisi.
    kind = fields.Selection(
        [('line', "Dépenses"), ('header', "En-tête")],
        string="Type", compute='_compute_kind', store=True)
    value = fields.Selection(
        VALUES, string="Valeur",
        help="Les totaux valent pour la ligne en colonne, pour toute la fiche "
             "en en-tête. Prestation, mois, période et société ne vont qu'en "
             "en-tête.")

    @api.depends('cell')
    def _compute_kind(self):
        for column in self:
            match = CELL_RE.match((column.cell or '').strip().upper())
            column.kind = 'header' if match and match.group(2) else 'line'

    @api.constrains('cell', 'value')
    def _check_cell(self):
        line_keys = {key for key, _label in LINE_VALUES}
        header_keys = {key for key, _label in HEADER_VALUES}
        for column in self:
            if not CELL_RE.match((column.cell or '').strip().upper()):
                raise UserError(_(
                    "« %(cell)s » : une lettre de colonne (B) ou une cellule (K4).",
                    cell=column.cell))
            allowed = header_keys if column.kind == 'header' else line_keys
            if not column.value:
                raise UserError(_("Choisissez la valeur à écrire en %s.", column.cell))
            if column.value not in allowed:
                raise UserError(_(
                    "« %(value)s » ne peut pas aller en %(where)s (%(cell)s).",
                    value=dict(VALUES)[column.value], cell=column.cell,
                    where=_("en-tête") if column.kind == 'header' else _("colonne de dépenses")))


class ExpenseScanSheet(models.AbstractModel):
    """Calcul des lignes, et fabrication des fichiers."""
    _name = 'expense.scan.sheet'
    _description = "Fabrication des fiches de frais"

    # ------------------------------------------------------------------
    # Lignes
    # ------------------------------------------------------------------

    @api.model
    def _is_flat_rate(self, expense):
        """Forfait ou kilométrage : pas de justificatif attendu."""
        product = expense.product_id
        return bool(product.standard_price) or expense._expense_scan_is_distance()

    @api.model
    def _receipts(self, expense):
        """Justificatifs de la dépense, le principal en premier."""
        if self._is_flat_rate(expense):
            return self.env['ir.attachment']
        attachments = expense.attachment_ids.filtered(
            lambda a: (a.mimetype or '').startswith('image/')
            or a.mimetype == 'application/pdf')
        main = expense.message_main_attachment_id & attachments
        return main | (attachments - main).sorted('id')

    @api.model
    def _tax_rate(self, expense):
        """Taux de TVA en fraction (0,1), texte s'il y en a plusieurs."""
        if expense._expense_scan_product_no_vat(expense.product_id) \
                or not expense.tax_amount:
            return 0.0
        if (expense.scan_detected_tax or '').startswith("plusieurs taux"):
            return "plusieurs taux"
        rates = expense.tax_ids.filtered(
            lambda t: t.amount_type == 'percent').mapped('amount')
        return rates[0] / 100.0 if len(rates) == 1 else ''

    @api.model
    def _lines(self, expenses):
        """Lignes numérotées, dans l'ordre chronologique."""
        ordered = expenses.sorted(lambda e: (
            e.date or date.min, e.scan_datetime or fields.Datetime.from_string('1970-01-01'), e.id))
        lines = []
        for number, expense in enumerate(ordered, start=1):
            receipts = self._receipts(expense)
            if len(receipts) > 1:
                labels = ["%s.%s" % (number, index) for index in range(1, len(receipts) + 1)]
                label = _("%(first)s à %(last)s", first=labels[0], last=labels[-1])
            else:
                labels = [str(number)] * len(receipts)
                label = str(number)
            if expense.expense_scan_nights_required:
                quantity = expense.expense_scan_nights or 1
            elif self._is_flat_rate(expense):
                quantity = expense.quantity or 1
            else:
                quantity = 1
            rate = self._tax_rate(expense)
            total_ttc = expense.total_amount
            # Un taux nul veut dire « pas de TVA récupérable » : forfait,
            # hôtel, transport, ou ticket sans TVA lisible.
            total_tva = 0.0 if rate == 0.0 else expense.tax_amount
            total_ht = total_ttc - total_tva
            lines.append({
                'expense': expense,
                'number': number,
                'n': label,
                'receipts': list(zip(labels, receipts)),
                'flat_rate': self._is_flat_rate(expense),
                'date': expense.date,
                'categorie': expense.product_id.name or '',
                'description': expense.name or '',
                'enseigne': expense.expense_scan_merchant or '',
                'mission': expense.project_id.name or '',
                'quantite': quantity,
                'prix_unitaire_ht': total_ht / quantity,
                'total_ht': total_ht,
                'taux_tva': rate,
                'taux_label': ("%g %%" % (rate * 100)) if isinstance(rate, float) else rate,
                'quantite_label': "%g" % quantity,
                'tva_unitaire': total_tva / quantity,
                'total_tva': total_tva,
                'ttc_unitaire': total_ttc / quantity,
                'total_ttc': total_ttc,
                'mode_paiement': dict(expense._fields['payment_mode']._description_selection(
                    self.env)).get(expense.payment_mode, ''),
                'salarie': expense.employee_id.name or '',
            })
        return lines

    @api.model
    def _header(self, expenses, lines):
        dates = [line['date'] for line in lines if line['date']]
        start, end = (min(dates), max(dates)) if dates else (False, False)
        # La mission n'est citée que si l'on a filtré dessus : des frais
        # sélectionnés à la main ne forment pas une prestation.
        projects = self.env['project.project'].browse(
            self.env.context.get('expense_scan_sheet_project_ids') or [])
        return {
            'salarie': ", ".join(expenses.employee_id.mapped('name')),
            'prestation': (", ".join(projects.mapped('name')) if projects
                           else _("Édition des frais sélectionnés")),
            'vat_rows': self._vat_summary(lines),
            'mois': start.replace(day=1) if start else False,
            'periode': (_("du %(start)s au %(end)s",
                          start=format_date(self.env, start), end=format_date(self.env, end))
                        if start else ''),
            'date_debut': start,
            'date_fin': end,
            'total_ht': sum(line['total_ht'] for line in lines),
            'total_tva': sum(line['total_tva'] for line in lines),
            'total_ttc': sum(line['total_ttc'] for line in lines),
            'societe': expenses.company_id[:1].name or '',
        }

    @api.model
    def _vat_summary(self, lines):
        """TVA récupérable par taux, pour la déclaration.

        Un taux vaut pour tous ses usages — 10 % de services ou de biens
        s'additionnent. Un ticket à taux mêlés, dont le détail n'est pas
        conservé, est compté à 20 %, comme une TVA dont le taux est inconnu.
        """
        groups = {}
        for line in lines:
            rate = line['taux_tva']
            mixed = not isinstance(rate, float)
            if mixed:
                rate = 0.2 if line['total_tva'] else 0.0
            group = groups.setdefault(rate, {
                'rate': rate, 'total_ht': 0.0, 'total_tva': 0.0, 'total_ttc': 0.0,
                'count': 0, 'mixed': False})
            group['total_ht'] += line['total_ht']
            group['total_tva'] += line['total_tva']
            group['total_ttc'] += line['total_ttc']
            group['count'] += 1
            group['mixed'] = group['mixed'] or (mixed and bool(line['total_tva']))
        rows = sorted(groups.values(), key=lambda group: -group['rate'])
        for row in rows:
            if row['rate']:
                row['label'] = "%g %%" % (row['rate'] * 100)
                if row['mixed']:
                    row['label'] += _(" (dont taux mêlés)")
            else:
                row['label'] = _("Sans TVA récupérable")
        return rows

    # ------------------------------------------------------------------
    # Contrôles propres au modèle
    # ------------------------------------------------------------------

    @api.model
    def _check_template(self, template, expenses):
        problems = []
        if template.reinvoice_only:
            wrong = expenses.filtered(lambda e: e.reinvoice_mode != 'project')
            if wrong:
                problems.append(_(
                    "Le modèle « %(template)s » n'accepte que des frais refacturables. "
                    "Ne le sont pas : %(names)s.", template=template.name,
                    names=", ".join("%s (%s)" % (e.name, format_date(self.env, e.date))
                                    for e in wrong)))
        if template.single_project and len(expenses.project_id) > 1:
            problems.append(_(
                "Le modèle « %(template)s » attend une seule mission ; la sélection "
                "en compte %(count)s : %(names)s.", template=template.name,
                count=len(expenses.project_id),
                names=", ".join(expenses.project_id.mapped('name'))))
        if problems:
            raise UserError("\n\n".join(problems))

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------

    @api.model
    def _excel(self, template, expenses):
        """Le classeur du modèle, rempli avec les dépenses d'un salarié."""
        try:
            import openpyxl  # noqa: PLC0415
            from openpyxl.utils import column_index_from_string  # noqa: PLC0415
            from openpyxl.formula.translate import Translator  # noqa: PLC0415
        except ImportError as error:
            raise UserError(_("La bibliothèque openpyxl manque sur le serveur.")) from error
        if not template.file:
            raise UserError(_("Le modèle « %s » n'a pas de fichier Excel.", template.name))
        self._check_template(template, expenses)

        lines = self._lines(expenses)
        header = self._header(expenses, lines)
        book = openpyxl.load_workbook(io.BytesIO(base64.b64decode(template.file)))
        sheet = book[template.sheet_name] if template.sheet_name else book.worksheets[0]

        first, last = template.first_row, template.last_row
        # Références prises avant tout déplacement : la première ligne donne
        # la mise en forme du corps du tableau, la dernière celle de son bas
        # (bordure de fin) et les formules d'une ligne vierge.
        body = {c.column: copy(c._style) for c in sheet[first] if c.has_style}
        bottom = {c.column: copy(c._style) for c in sheet[last] if c.has_style}
        formulas = {c.column: c.value for c in sheet[last]
                    if isinstance(c.value, str) and c.value.startswith('=')}
        height = sheet.row_dimensions[first].height

        # Exactement une ligne par dépense : les lignes en trop disparaissent,
        # celles qui manquent sont ajoutées, et les totaux suivent.
        self._resize_table(sheet, last, len(lines) - (last - first + 1))
        new_last = first + len(lines) - 1

        columns = {column_index_from_string(c.cell.strip().upper()): c.value
                   for c in template.column_ids if c.kind == 'line' and c.value}

        for index, line in enumerate(lines):
            row = first + index
            dimension = sheet.row_dimensions[row]
            dimension.hidden = False
            dimension.height = height
            for column, style in body.items():
                sheet.cell(row=row, column=column)._style = copy(style)
            for column, formula in formulas.items():
                if column not in columns:
                    sheet.cell(row=row, column=column).value = Translator(
                        formula, origin=sheet.cell(row=last, column=column).coordinate
                    ).translate_formula(sheet.cell(row=row, column=column).coordinate)
            for column, value in columns.items():
                sheet.cell(row=row, column=column).value = line[value] if line[value] != '' else None
        for column, style in bottom.items():
            sheet.cell(row=new_last, column=column)._style = copy(style)

        for column in template.column_ids.filtered(lambda c: c.kind == 'header' and c.value):
            sheet[column.cell.strip().upper()].value = header[column.value] or None

        output = io.BytesIO()
        book.save(output)
        return output.getvalue()

    @api.model
    def _resize_table(self, sheet, last, delta):
        """Ajoute (``delta`` > 0) ou retire (< 0) des lignes en fin de tableau.

        openpyxl déplace les cellules, mais ni les hauteurs de ligne, ni
        l'état masqué, ni les fusions, ni les plages des formules situées
        plus bas : tout cela est recalé ici, pour que la ligne des totaux
        garde son aspect et somme exactement les lignes du tableau.
        """
        if not delta:
            return
        new_last = last + delta
        below = {row: (dim.height, dim.hidden)
                 for row, dim in sheet.row_dimensions.items() if row > last}
        if delta > 0:
            sheet.insert_rows(last + 1, delta)
        else:
            sheet.delete_rows(new_last + 1, -delta)
        for row in [row for row in sheet.row_dimensions if row > new_last]:
            del sheet.row_dimensions[row]
        for row, (height, hidden) in below.items():
            dimension = sheet.row_dimensions[row + delta]
            dimension.height = height
            dimension.hidden = hidden

        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row > last:
                merged.shift(0, delta)
            elif merged.max_row > new_last:
                # Fusion prise dans les lignes retirées : elle n'a plus lieu d'être.
                sheet.merged_cells.remove(merged)
        def shift(match):
            start = int(match.group(2))
            text = "%s%s" % (match.group(1), start + delta if start > last else start)
            if match.group(3):
                end = int(match.group(4))
                if end > last:
                    end += delta
                elif end == last and start <= last:
                    end = new_last  # plage qui court jusqu'au bas du tableau
                text += ":%s%s" % (match.group(3), end)
            return text

        # Toute la feuille : un total peut être repris ailleurs, en en-tête
        # (« =L70 ») ou dans une autre formule du pied de tableau.
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.value = CELL_REF_RE.sub(shift, cell.value)

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    @api.model
    def _summary_pdf(self, expenses, with_receipts=True):
        report = self.env.ref('expense_scan.action_report_expense_sheet')
        lines = self._lines(expenses)
        pdf, _kind = self.env['ir.actions.report']._render_qweb_pdf(
            report, res_ids=[line['expense'].id for line in lines],
            data={'expense_ids': expenses.ids})
        if not with_receipts:
            return pdf
        receipts = self._receipts_pdf(lines)
        if not receipts:
            return pdf
        from odoo.tools.pdf import merge_pdf  # noqa: PLC0415
        return merge_pdf([pdf, receipts])

    @api.model
    def _receipts_pdf(self, lines):
        """Les justificatifs, un par page, chacun marqué de son numéro."""
        from odoo.tools.pdf import PdfFileReader, PdfFileWriter  # noqa: PLC0415
        writer = PdfFileWriter()
        pages = 0
        for line in lines:
            for label, attachment in line['receipts']:
                for page in self._attachment_pages(attachment, label, PdfFileReader):
                    writer.addPage(page)
                    pages += 1
        if not pages:
            return b''
        output = io.BytesIO()
        writer.write(output)
        return output.getvalue()

    @api.model
    def _attachment_pages(self, attachment, label, PdfFileReader):
        raw = attachment.raw or b''
        mimetype = attachment.mimetype or ''
        try:
            if mimetype.startswith('image/'):
                stream = io.BytesIO(self._image_page(raw, label))
                return [PdfFileReader(stream, strict=False).getPage(0)]
            reader = PdfFileReader(io.BytesIO(raw), strict=False)
            count = reader.getNumPages()
            pages = []
            for index in range(count):
                page = reader.getPage(index)
                width = float(abs(page.mediaBox.getWidth()))
                height = float(abs(page.mediaBox.getHeight()))
                text = label if count == 1 else "%s (%s/%s)" % (label, index + 1, count)
                stamp = PdfFileReader(io.BytesIO(self._stamp(text, width, height)), strict=False)
                page.mergePage(stamp.getPage(0))
                pages.append(page)
            return pages
        except Exception:  # noqa: BLE001 - justificatif illisible
            _logger.warning("Justificatif illisible : %s", attachment.name, exc_info=True)
            stream = io.BytesIO(self._unreadable_page(attachment.name, label))
            return [PdfFileReader(stream, strict=False).getPage(0)]

    @api.model
    def _image_page(self, raw, label):
        """Une page A4 portant la photo, à la plus grande taille possible."""
        from PIL import Image, ImageOps  # noqa: PLC0415
        from reportlab.lib.pagesizes import A4  # noqa: PLC0415
        from reportlab.lib.utils import ImageReader  # noqa: PLC0415
        from reportlab.pdfgen import canvas  # noqa: PLC0415

        image = ImageOps.exif_transpose(Image.open(io.BytesIO(raw))).convert('RGB')
        # Une photo de téléphone pèse plusieurs Mo : réduite à une taille
        # lisible à l'impression et passée en JPEG, elle est incorporée telle
        # quelle au PDF au lieu d'y être stockée pixel par pixel.
        image.thumbnail((IMAGE_MAX_SIDE, IMAGE_MAX_SIDE))
        jpeg = io.BytesIO()
        image.save(jpeg, format='JPEG', quality=IMAGE_QUALITY, optimize=True)
        jpeg.seek(0)
        width, height = A4
        margin, band = 28, 40
        box_w, box_h = width - 2 * margin, height - 2 * margin - band
        ratio = min(box_w / image.width, box_h / image.height)
        draw_w, draw_h = image.width * ratio, image.height * ratio
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        self._draw_label(pdf, label, width, height)
        pdf.drawImage(ImageReader(jpeg), (width - draw_w) / 2,
                      margin + (box_h - draw_h) / 2, draw_w, draw_h)
        pdf.showPage()
        pdf.save()
        return output.getvalue()

    @api.model
    def _stamp(self, label, width, height):
        from reportlab.pdfgen import canvas  # noqa: PLC0415
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=(width, height))
        self._draw_label(pdf, label, width, height)
        pdf.showPage()
        pdf.save()
        return output.getvalue()

    @api.model
    def _unreadable_page(self, name, label):
        from reportlab.lib.pagesizes import A4  # noqa: PLC0415
        from reportlab.pdfgen import canvas  # noqa: PLC0415
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        self._draw_label(pdf, label, *A4)
        pdf.setFont('Helvetica', 12)
        pdf.drawString(40, A4[1] / 2, _("Justificatif illisible : %s", name or ''))
        pdf.showPage()
        pdf.save()
        return output.getvalue()

    @api.model
    def _draw_label(self, pdf, label, width, height):
        """« N° 5.2 » dans un cartouche blanc, en haut à droite."""
        text = _("N° %s", label)
        pdf.setFont('Helvetica-Bold', 16)
        text_w = pdf.stringWidth(text, 'Helvetica-Bold', 16)
        pdf.setFillColorRGB(1, 1, 1)
        pdf.setStrokeColorRGB(0, 0, 0)
        pdf.rect(width - text_w - 36, height - 38, text_w + 20, 26, fill=1, stroke=1)
        pdf.setFillColorRGB(0, 0, 0)
        pdf.drawString(width - text_w - 26, height - 30, text)

    # ------------------------------------------------------------------
    # Assemblage
    # ------------------------------------------------------------------

    @api.model
    def _file_stem(self, expenses, prefix):
        lines = self._lines(expenses)
        header = self._header(expenses, lines)
        start = header['date_debut']
        period = start.strftime('%Y-%m') if start else ''
        stem = " - ".join(part for part in (prefix, header['salarie'], period) if part)
        return re.sub(r'[\\/:*?"<>|]+', '-', stem)

    @api.model
    def _build(self, expenses, summary=False, excel_template=False, receipts=False):
        """``[(nom de fichier, contenu)]`` pour chaque salarié et chaque sortie."""
        files = []
        for employee in expenses.employee_id:
            own = expenses.filtered(lambda e: e.employee_id == employee)
            if summary:
                files.append(("%s.pdf" % self._file_stem(own, _("Fiche de frais")),
                              self._summary_pdf(own)))
            if excel_template:
                files.append(("%s.xlsx" % self._file_stem(own, excel_template.name),
                              self._excel(excel_template, own)))
            if receipts:
                content = self._receipts_pdf(self._lines(own))
                if content:
                    files.append(("%s.pdf" % self._file_stem(own, _("Justificatifs")), content))
        return files

    @api.model
    def _zip(self, files):
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
            seen = set()
            for name, content in files:
                unique, index = name, 2
                while unique in seen:
                    stem, dot, ext = name.rpartition('.')
                    unique = "%s (%s).%s" % (stem, index, ext)
                    index += 1
                seen.add(unique)
                archive.writestr(unique, content)
        return output.getvalue()


class ExpenseSheetReport(models.AbstractModel):
    _name = 'report.expense_scan.report_expense_sheet'
    _description = "Fiche de frais (PDF)"

    @api.model
    def _get_report_values(self, docids, data=None):
        ids = (data or {}).get('expense_ids') or docids
        expenses = self.env['hr.expense'].browse(ids)
        Sheet = self.env['expense.scan.sheet']
        lines = Sheet._lines(expenses)
        header = Sheet._header(expenses, lines)
        currency = expenses.company_id[:1].currency_id or self.env.company.currency_id
        return {
            'doc_ids': expenses.ids,
            'doc_model': 'hr.expense',
            'docs': expenses,
            'o': expenses[:1],
            'company': expenses.company_id[:1] or self.env.company,
            'lines': lines,
            'header': header,
            'currency': currency,
            'format_date': lambda value: format_date(self.env, value) if value else '',
        }


class ExpenseScanSheetWizard(models.TransientModel):
    _name = 'expense.scan.sheet.wizard'
    _description = "Édition d'une fiche de frais"

    expense_ids = fields.Many2many('hr.expense', string="Dépenses sélectionnées")
    scope = fields.Selection(
        [('all', "Tous les frais"), ('reinvoice', "Frais refacturables seulement")],
        string="Frais", default='reinvoice', required=True)
    project_ids = fields.Many2many(
        'project.project', string="Missions",
        help="Vide : toutes les missions de la sélection.")
    available_project_ids = fields.Many2many(
        'project.project', compute='_compute_available_project_ids')
    summary = fields.Boolean(string="Récapitulatif PDF avec justificatifs", default=True)
    excel = fields.Boolean(string="Export Excel")
    template_id = fields.Many2one('expense.scan.export.template', string="Modèle")
    receipts = fields.Boolean(string="Justificatifs seuls (PDF)")
    selected_count = fields.Integer(compute='_compute_selected')
    selected_summary = fields.Char(compute='_compute_selected')
    result_file = fields.Binary(readonly=True, attachment=False)
    result_name = fields.Char(readonly=True)

    @api.model
    def default_get(self, fields_list):
        values = super().default_get(fields_list)
        if self.env.context.get('active_model') == 'hr.expense':
            values['expense_ids'] = [(6, 0, self.env.context.get('active_ids') or [])]
        return values

    @api.depends('expense_ids')
    def _compute_available_project_ids(self):
        for wizard in self:
            wizard.available_project_ids = wizard.expense_ids.project_id

    @api.onchange('excel')
    def _onchange_excel(self):
        # Le premier modèle de la liste — son ordre se règle dans la
        # configuration — est proposé d'office.
        if self.excel and not self.template_id:
            self.template_id = self.env['expense.scan.export.template'].search([], limit=1)

    @api.onchange('template_id')
    def _onchange_template_id(self):
        # Un modèle réservé aux frais refacturables règle le filtre d'office.
        if self.template_id.reinvoice_only:
            self.scope = 'reinvoice'

    def _selected(self):
        self.ensure_one()
        # Pendant l'édition du formulaire, les enregistrements liés sont des
        # copies provisoires : on compare les vrais, sans quoi une mission
        # choisie ne correspondait à aucune dépense (« 0 dépense »).
        expenses = self.expense_ids._origin
        if self.scope == 'reinvoice':
            expenses = expenses.filtered(lambda e: e.reinvoice_mode == 'project')
        project_ids = set(self.project_ids._origin.ids)
        if project_ids:
            expenses = expenses.filtered(lambda e: e.project_id.id in project_ids)
        return expenses

    @api.depends('expense_ids', 'scope', 'project_ids')
    def _compute_selected(self):
        for wizard in self:
            selected = wizard._selected()
            wizard.selected_count = len(selected)
            wizard.selected_summary = _(
                "%(count)s dépense(s), %(employees)s salarié(s), %(amount).2f TTC",
                count=len(selected), employees=len(selected.employee_id),
                amount=sum(selected.mapped('total_amount')))

    def action_generate(self):
        self.ensure_one()
        expenses = self._selected()
        if not expenses:
            raise UserError(_("Aucune dépense ne correspond aux filtres."))
        if not (self.summary or self.excel or self.receipts):
            raise UserError(_("Choisissez au moins une sortie."))
        if self.excel and not self.template_id:
            raise UserError(_("Choisissez le modèle Excel."))
        Sheet = self.env['expense.scan.sheet'].with_context(
            expense_scan_sheet_project_ids=self.project_ids.ids)
        expenses = expenses.with_context(expense_scan_sheet_project_ids=self.project_ids.ids)
        files = Sheet._build(
            expenses, summary=self.summary,
            excel_template=self.excel and self.template_id,
            receipts=self.receipts)
        if not files:
            raise UserError(_("Rien à produire : aucune de ces dépenses n'a de justificatif."))
        if len(files) == 1:
            name, content = files[0]
        else:
            name = "%s.zip" % self.env['expense.scan.sheet']._file_stem(expenses, _("Fiches de frais"))
            content = self.env['expense.scan.sheet']._zip(files)
        self.write({'result_file': base64.b64encode(content), 'result_name': name})
        # Téléchargé par le client, qui ferme ensuite la fenêtre : un simple
        # lien la laissait ouverte, et un nouvel onglet serait bloqué.
        return {
            'type': 'ir.actions.client',
            'tag': 'expense_scan_download',
            'params': {
                'url': '/web/content/?model=%s&id=%s&field=result_file'
                       '&filename_field=result_name&download=true' % (self._name, self.id),
            },
        }
